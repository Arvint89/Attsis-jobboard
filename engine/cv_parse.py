"""
cv_parse.py -- CV input + parsing (Phase 1, C1+C2).

Reads a CV file (.md/.txt/.docx) and writes profile.json: the matching model the
scorer runs on. It starts from model_base.BASE (curated titles + exclusions +
booster domains BT can ramp into), then sets skill_confirmers to the subset the
CV actually backs -- so a different CV produces a different scoreboard. It also
records cv_evidence (which model terms are/aren't supported by the CV), which is
the hook Phase 2's gap analysis and the project's no-overclaim rule will use.

    python cv_parse.py path/to/CV.docx          -> writes engine/profile.json
    python cv_parse.py --cv CV.md --out profile.json

No LLM. .docx needs python-docx (pip install python-docx); .md/.txt need nothing.
"""
from __future__ import annotations
import argparse
import json
import os
import re
import sys

import model_base
import geo


def extract_person(text: str):
    for line in text.splitlines():
        line = line.strip()
        if not line or any(ch.isdigit() for ch in line):
            continue
        # strip trailing credentials (", P.Eng.", ", PhD")
        core = re.split(r",|\|", line)[0].strip()
        words = core.split()
        if 2 <= len(words) <= 4 and all(w[:1].isupper() and w.replace('.', '').isalpha() for w in words):
            initials = "".join(w[0] for w in words if w[0].isupper()).upper()
            return core, initials
    return "", ""


def extract_home(text: str):
    low = text.lower()
    for city in geo.CITY_COORDS:
        if city and city != "remote" and city in low:
            return city.title()
    return ""

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
MIN_CONFIRMERS = 8   # safety floor: below this, keep the full base set


def read_cv(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt", ""):
        return open(path, encoding="utf-8", errors="ignore").read()
    if ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            raise SystemExit("python-docx needed for .docx: pip install python-docx")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    raise SystemExit(f"unsupported CV type: {ext} (use .md/.txt/.docx)")


def extract_years(text: str):
    m = re.search(r"(\d{1,2})\s*\+?\s*years", text.lower())
    return int(m.group(1)) if m else None


def build_profile(cv_text: str, source: str, name=None, home=None):
    low = cv_text.lower()
    base = model_base.BASE
    # which base confirmers does the CV back?
    backed = [c for c in base["skill_confirmers"] if c in low]
    confirmers = sorted(set(backed)) if len(backed) >= MIN_CONFIRMERS else list(base["skill_confirmers"])

    # evidence report across the whole model vocabulary
    def ev(terms):
        return {"backed": [t for t in terms if t in low],
                "unbacked": [t for t in terms if t not in low]}

    cv_evidence = {
        "confirmers": ev(base["skill_confirmers"]),
        "medical": ev(base["boosters"]["medical"]),
        "iec60601": ev(base["boosters"]["iec60601"]),
        "space": ev(base["boosters"]["space"]),
        "verification": ev(base["boosters"]["verification"]),
        "title_core": ev(base["title_triggers"]["core"]),
    }

    model = {
        "title_triggers": base["title_triggers"],   # curated
        "skill_confirmers": confirmers,              # CV-DRIVEN
        "boosters": base["boosters"],                # domains to ramp into (discovery)
        "exclusions": base["exclusions"],            # curated
        "params": base["params"],
    }
    person, initials = extract_person(cv_text)
    return {
        "generated_from": os.path.basename(source),
        "person": name or person,
        "initials": (name[:2].upper() if name else initials) or "?",
        "home": home or extract_home(cv_text),
        "years_general": extract_years(cv_text),
        "confirmers_backed_by_cv": len(backed),
        "confirmers_total_base": len(base["skill_confirmers"]),
        "model": model,
        "cv_evidence": cv_evidence,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cv", nargs="?", help="path to CV (.md/.txt/.docx)")
    ap.add_argument("--cv", dest="cv_opt")
    ap.add_argument("--out", default=os.path.join(HERE, "profile.json"))
    ap.add_argument("--name")
    ap.add_argument("--home", help="home city for distance rings (e.g. London)")
    a = ap.parse_args()
    path = a.cv or a.cv_opt
    if not path:
        raise SystemExit("give a CV path: python cv_parse.py CV.docx")
    text = read_cv(path)
    prof = build_profile(text, path, name=a.name, home=a.home)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(prof, f, indent=2, ensure_ascii=False)
    print(f"profile <- {os.path.basename(path)}: person={prof['person'] or '?'} "
          f"({prof['initials']}), home={prof['home'] or '?'}, years={prof['years_general']}, "
          f"{len(prof['model']['skill_confirmers'])} active confirmers -> {a.out}")


if __name__ == "__main__":
    main()
